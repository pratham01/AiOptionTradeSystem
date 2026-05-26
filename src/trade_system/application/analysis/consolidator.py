import pandas as pd
import logging
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

logger = logging.getLogger(__name__)

class DataConsolidator:
    def __init__(self, data_dir: Path):
        self.data_dir = data_dir

    def to_long_format(self, output_file: Path) -> pd.DataFrame:
        """Consolidates individual stock CSVs into a single master Long-Format file."""
        all_files = list(self.data_dir.glob("*.csv"))
        if not all_files:
            logger.error(f"No CSV files found in {self.data_dir}")
            return pd.DataFrame()

        logger.info(f"Consolidating {len(all_files)} files into Long Format...")
        all_dfs = []
        for filename in all_files:
            try:
                # Extract symbol (assumes format like NSE_SYMBOL-EQ_Daily.csv)
                symbol = filename.name.split('_')[1].replace('-EQ', '').replace('-INDEX', '')
                df = pd.read_csv(filename)
                if df.empty: continue
                df['symbol'] = symbol
                all_dfs.append(df)
            except Exception as e:
                logger.error(f"Error processing {filename}: {e}")

        if not all_dfs: return pd.DataFrame()

        master_df = pd.concat(all_dfs, ignore_index=True)
        # Reorder and sort
        cols = ['timestamp', 'symbol', 'open', 'high', 'low', 'close', 'volume']
        master_df = master_df[[c for c in cols if c in master_df.columns]]
        master_df = master_df.sort_values(['timestamp', 'symbol'])
        
        master_df.to_csv(output_file, index=False)
        logger.info(f"Saved Long-Format master to {output_file}")
        return master_df

    def to_formatted_excel(self, output_file: Path):
        """Creates a professionally formatted Wide-Format Excel file with merged headers."""
        all_files = list(self.data_dir.glob("*.csv"))
        if not all_files: return

        logger.info(f"Building Formatted Excel Master from {len(all_files)} files...")
        symbol_data = {}
        all_timestamps = set()
        
        for filename in all_files:
            try:
                symbol = filename.name.split('_')[1].replace('-EQ', '').replace('-INDEX', '')
                df = pd.read_csv(filename, parse_dates=['timestamp'])
                if df.empty: continue
                df = df.set_index('timestamp')
                symbol_data[symbol] = df[['open', 'high', 'low', 'close', 'volume']]
                all_timestamps.update(df.index.tolist())
            except Exception as e:
                logger.error(f"Error reading {filename}: {e}")

        sorted_ts = sorted(list(all_timestamps))
        wb = Workbook()
        ws = wb.active
        ws.title = "F&O Historical Master"

        # Styles
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        sub_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        alignment_center = Alignment(horizontal="center", vertical="center")
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        # Date column
        ws.cell(row=1, column=1, value="Date").alignment = alignment_center
        ws.merge_cells(start_row=1, end_row=2, start_column=1, end_column=1)
        
        col_idx = 2
        sorted_symbols = sorted(symbol_data.keys())
        for symbol in sorted_symbols:
            # Merged Symbol Header
            cell = ws.cell(row=1, column=col_idx, value=symbol)
            cell.alignment = alignment_center
            cell.font = header_font
            cell.fill = header_fill
            ws.merge_cells(start_row=1, end_row=1, start_column=col_idx, end_column=col_idx+4)
            
            # Sub-headers
            sub_headers = ['Open', 'High', 'Low', 'Close', 'Volume']
            for i, sh in enumerate(sub_headers):
                sub_cell = ws.cell(row=2, column=col_idx + i, value=sh)
                sub_cell.alignment = alignment_center
                sub_cell.fill = sub_header_fill
                sub_cell.border = border
            col_idx += 5

        # Data Rows
        row_idx = 3
        for ts in sorted_ts:
            ws.cell(row=row_idx, column=1, value=ts.strftime("%Y-%m-%d"))
            col_idx = 2
            for symbol in sorted_symbols:
                data = symbol_data[symbol]
                if ts in data.index:
                    row = data.loc[ts]
                    ws.cell(row=row_idx, column=col_idx, value=row['open'])
                    ws.cell(row=row_idx, column=col_idx+1, value=row['high'])
                    ws.cell(row=row_idx, column=col_idx+2, value=row['low'])
                    ws.cell(row=row_idx, column=col_idx+3, value=row['close'])
                    ws.cell(row=row_idx, column=col_idx+4, value=row['volume'])
                col_idx += 5
            row_idx += 1

        ws.freeze_panes = "B3"
        wb.save(output_file)
        logger.info(f"Saved formatted Excel master to {output_file}")
